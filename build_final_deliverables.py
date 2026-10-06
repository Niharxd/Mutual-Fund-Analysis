"""Build the capstone PDF report and 12-slide presentation from project outputs."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as PDFImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "reports"
CHARTS = REPORTS / "charts"
PAGES = REPORTS / "dashboard_pages"
NAVY = "14233B"
BLUE = "2457A7"
TEAL = "27A6A1"
PALE = "F2F6FC"
WHITE = "FFFFFF"
MUTED = "5C6878"


def _read_report(name: str) -> pd.DataFrame:
    """Read a generated report CSV and fail clearly when a prerequisite is missing."""
    path = REPORTS / name
    if not path.exists():
        raise FileNotFoundError(f"Required report output is missing: {path}")
    return pd.read_csv(path)


def _read_data(name: str) -> pd.DataFrame:
    """Read a processed dataset used for a headline statistic."""
    path = ROOT / "data" / "processed" / name
    if not path.exists():
        raise FileNotFoundError(f"Required processed dataset is missing: {path}")
    return pd.read_csv(path)


def _format_percent(value: object, digits: int = 1) -> str:
    """Format a decimal ratio as a percentage, preserving missing values."""
    number = pd.to_numeric(value, errors="coerce")
    return "n/a" if pd.isna(number) else f"{number:.{digits}%}"


def _chart(name: str, width: float = 7.0, height: float = 2.6) -> PDFImage:
    """Create a proportionally scaled report image and validate its path."""
    path = CHARTS / name
    if not path.exists():
        path = REPORTS / name
    if not path.exists():
        path = PAGES / name
    if not path.exists():
        raise FileNotFoundError(f"Required report image is missing: {path}")
    image = PDFImage(str(path))
    scale = min(width * inch / image.imageWidth, height * inch / image.imageHeight)
    image.drawWidth = image.imageWidth * scale
    image.drawHeight = image.imageHeight * scale
    image.hAlign = "CENTER"
    return image


def _pdf_table(rows: list[list[object]], widths: list[float] | None = None) -> Table:
    """Style a compact table for the report."""
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2457A7")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("LEADING", (0, 0), (-1, -1), 10),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FC")]),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E2EF")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _pdf_styles() -> dict[str, ParagraphStyle]:
    """Create consistent report typography."""
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "ReportTitle", parent=base["Title"], fontName="Helvetica-Bold",
            fontSize=27, leading=32, textColor=colors.HexColor(f"#{NAVY}"),
            alignment=TA_LEFT, spaceAfter=14,
        ),
        "heading": ParagraphStyle(
            "SectionHeading", parent=base["Heading1"], fontName="Helvetica-Bold",
            fontSize=19, leading=23, textColor=colors.HexColor(f"#{NAVY}"),
            spaceAfter=12,
        ),
        "subheading": ParagraphStyle(
            "Subheading", parent=base["Heading2"], fontName="Helvetica-Bold",
            fontSize=11, leading=14, textColor=colors.HexColor(f"#{BLUE}"),
            spaceBefore=8, spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "ReportBody", parent=base["BodyText"], fontName="Helvetica",
            fontSize=9.4, leading=14, textColor=colors.HexColor("#344054"),
            spaceAfter=8,
        ),
        "small": ParagraphStyle(
            "ReportSmall", parent=base["BodyText"], fontName="Helvetica",
            fontSize=8, leading=11, textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=6,
        ),
        "cover": ParagraphStyle(
            "CoverText", parent=base["BodyText"], fontName="Helvetica",
            fontSize=13, leading=19, textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=12,
        ),
        "center": ParagraphStyle(
            "Centered", parent=base["BodyText"], fontName="Helvetica",
            fontSize=10, leading=14, alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{MUTED}"),
        ),
    }


def _paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    """Create a ReportLab paragraph from plain text."""
    return Paragraph(text, style)


def _footer(canvas: object, document: object) -> None:
    """Draw a consistent page number and footer line."""
    canvas.saveState()
    page_width, _ = letter
    canvas.setStrokeColor(colors.HexColor("#D9E2EF"))
    canvas.line(0.65 * inch, 0.52 * inch, page_width - 0.65 * inch, 0.52 * inch)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.drawString(0.65 * inch, 0.34 * inch, "Bluestock Mutual Fund Analysis | Capstone")
    canvas.drawRightString(page_width - 0.65 * inch, 0.34 * inch, f"Page {document.page}")
    canvas.restoreState()


def build_pdf() -> Path:
    """Generate a 16-page, source-grounded capstone report."""
    from analytics.advanced_analytics import sector_hhi, sip_continuity

    master = _read_data("01_fund_master.csv")
    nav = _read_data("02_nav_history.csv")
    aum = _read_data("03_aum_by_fund_house.csv")
    sip = _read_data("04_monthly_sip_inflows.csv")
    folios = _read_data("06_industry_folio_count.csv")
    transactions = _read_data("08_investor_transactions.csv")
    holdings = _read_data("09_portfolio_holdings.csv")
    transactions["transaction_date"] = pd.to_datetime(transactions["transaction_date"])
    performance = _read_report("fund_scorecard.csv")
    var_cvar = _read_report("var_cvar_report.csv")
    styles = _pdf_styles()
    body, small, heading, sub = styles["body"], styles["small"], styles["heading"], styles["subheading"]
    story: list[object] = []

    def page(title: str, content: Iterable[object]) -> None:
        """Append a titled report page and a page break."""
        story.append(_paragraph(title, heading))
        story.extend(content)
        story.append(PageBreak())

    # 1. Cover
    logo = ROOT / "dashboard" / "assets" / "bluestock_logo.png"
    if logo.exists():
        cover_logo = PDFImage(str(logo), width=1.75 * inch, height=1.05 * inch, kind="proportional")
        story.extend([Spacer(1, 0.85 * inch), cover_logo, Spacer(1, 0.38 * inch)])
    else:
        story.append(Spacer(1, 1.6 * inch))
    story.extend(
        [
            _paragraph("Bluestock Mutual Fund Analysis", styles["title"]),
            _paragraph("Final Capstone Report", ParagraphStyle(
                "CoverSubtitle", parent=styles["title"], fontSize=20, leading=25,
                textColor=colors.HexColor(f"#{BLUE}"),
            )),
            Spacer(1, 0.18 * inch),
            _paragraph(
                "Industry context, fund performance, risk, investor behaviour, "
                "and an interactive Streamlit dashboard.",
                styles["cover"],
            ),
            Spacer(1, 0.28 * inch),
            _pdf_table(
                [
                    ["Scope", "Evidence in this report"],
                    ["Datasets", "10 supplied CSV sources, transformed to a SQLite-backed analytical model"],
                    ["Fund sample", f"{master['amfi_code'].nunique()} schemes; {nav['amfi_code'].nunique()} with observed NAV"],
                    [
                        "Analysis window",
                        f"Daily NAV observations: {pd.to_datetime(nav['date']).min():%Y-%m-%d} "
                        f"to {pd.to_datetime(nav['date']).max():%Y-%m-%d}",
                    ],
                    ["Interactive output", "Four-page Streamlit dashboard; no hosted URL or native PBIX is claimed"],
                ],
                [1.25 * inch, 5.5 * inch],
            ),
            Spacer(1, 0.55 * inch),
            _paragraph("Prepared from the project datasets and generated analysis outputs.", styles["center"]),
            PageBreak(),
        ]
    )

    # 2. Executive summary
    latest_aum_date = pd.to_datetime(aum["date"]).max()
    latest_aum = aum.loc[pd.to_datetime(aum["date"]).eq(latest_aum_date), "aum_lakh_crore"].sum()
    latest_sip = float(sip.iloc[-1]["sip_inflow_crore"])
    latest_folios = float(folios.iloc[-1]["total_folios_crore"])
    top = performance.sort_values("score_rank").iloc[0]
    worst_var = var_cvar.sort_values("var_95_daily_return_pct").iloc[0]
    _, continuity_summary = sip_continuity(transactions)
    continuity_metrics = continuity_summary.iloc[0]
    eligible_sip_investors = int(continuity_metrics["eligible_investors_6plus_sips"])
    at_risk_sip_investors = int(continuity_metrics["at_risk_investors_avg_gap_over_35d"])
    at_risk_pct = (
        100 * at_risk_sip_investors / eligible_sip_investors
        if eligible_sip_investors
        else float("nan")
    )
    hhi = sector_hhi(holdings, master).dropna(subset=["hhi"])
    top_hhi = hhi.iloc[0] if not hhi.empty else None
    page(
        "1. Executive summary",
        [
            _paragraph(
                f"The project combines industry snapshots, 40-scheme NAV history, performance and "
                f"risk analytics, a 32,778-row investor transaction sample, and an interactive dashboard. "
                f"The latest supplied AUM snapshot ({latest_aum_date:%b %Y}) totals INR {latest_aum:.2f} "
                f"lakh crore across reported fund-house records. December 2025 SIP inflow is INR "
                f"{latest_sip:,.0f} crore; the latest folio snapshot records {latest_folios:.2f} crore.",
                body,
            ),
            _pdf_table(
                [
                    ["Headline", "Result", "Interpretation"],
                    ["Scorecard leader", str(top["scheme_name"]), f"Composite score {top['fund_score']:.2f}/100"],
                    ["Lowest 95% daily VaR", str(worst_var["scheme_name"]), f"{worst_var['var_95_daily_return_pct']:.2f}% threshold"],
                    ["Fund coverage", f"{master['amfi_code'].nunique()} schemes", "Sample, not a census of all Indian mutual funds"],
                    ["NAV observation dates", f"{pd.to_datetime(nav['date']).nunique():,}", "Market-date series used for return analysis"],
                ],
                [1.45 * inch, 2.5 * inch, 2.8 * inch],
            ),
            Spacer(1, 0.2 * inch),
            _paragraph(
                "The strongest practical outcome is a reproducible local analytics workflow: cleaned "
                "data, SQLite model, notebooks and scripts, downloadable static exports, and a Streamlit "
                "dashboard. Performance rankings are descriptive and should not be read as investment advice.",
                body,
            ),
        ],
    )

    # 3. Objectives and scope
    page(
        "2. Objectives and completion review",
        [
            _pdf_table(
                [
                    ["Objective", "Project evidence", "Status"],
                    ["Data ingestion and validation", "10 CSV inputs; code/NAV coverage checks; live NAV fetch utility", "Complete"],
                    ["Cleaning and analytical model", "Processed datasets, SQLite database, schema and query SQL", "Complete"],
                    ["EDA and performance", "EDA, performance and advanced analytics notebooks/reports", "Complete"],
                    ["Risk and investor analytics", "VaR/CVaR, rolling Sharpe, cohort/SIP continuity, HHI, recommender", "Complete"],
                    ["Interactive dashboard", "Four-page Streamlit app with filters and charts", "Complete"],
                    ["Native Power BI PBIX", "Generated PBIP was not validated in Desktop", "Not delivered"],
                    ["Public dashboard publishing", "No hosting/account destination configured", "Not published"],
                ],
                [1.7 * inch, 4.35 * inch, 0.75 * inch],
            ),
            Spacer(1, 0.18 * inch),
            _paragraph(
                "This review distinguishes implemented local deliverables from optional platform deliverables. "
                "The Streamlit app is the supported interactive route in this repository; neither a published "
                "URL nor a validated Power BI PBIX is represented as complete.",
                body,
            ),
        ],
    )

    # 4. Sources
    datasets = [
        ["Source file", "Purpose", "Observed coverage / shape"],
        ["01_fund_master.csv", "Scheme metadata", "40 rows"],
        ["02_nav_history.csv", "Daily scheme NAV", "46,000 rows; 40 schemes"],
        ["03_aum_by_fund_house.csv", "AMC-level AUM snapshots", "90 rows; 10 AMCs"],
        ["04_monthly_sip_inflows.csv", "Industry SIP trends", "48 months"],
        ["05_category_inflows.csv", "Category flows", "144 rows; FY25 coverage"],
        ["06_industry_folio_count.csv", "Folio trends", "21 monthly snapshots"],
        ["07_scheme_performance.csv", "Scheme performance attributes", "40 rows"],
        ["08_investor_transactions.csv", "Synthetic/sample transaction records", "32,778 rows; Jan 2024-May 2025"],
        ["09_portfolio_holdings.csv", "Scheme holdings and sector weights", "322 rows; 34 schemes"],
        ["10_benchmark_indices.csv", "Market/benchmark index levels", "8,050 rows"],
    ]
    page(
        "3. Data sources and coverage",
        [
            _paragraph(
                "The supplied local CSVs are the analytical source of truth. A live NAV helper also queries "
                "MFAPI for selected public scheme histories; the pipeline does not make external requests by default.",
                body,
            ),
            _pdf_table(datasets, [2.0 * inch, 2.05 * inch, 2.75 * inch]),
            Spacer(1, 0.15 * inch),
            _paragraph(
                "The datasets have different reporting windows and grains; dates were not artificially "
                "aligned for all analyses. Investor data is treated as a provided sample, not a representative "
                "population survey.",
                small,
            ),
        ],
    )

    # 5. Data quality
    code_count = master["amfi_code"].nunique()
    nav_codes = nav["amfi_code"].nunique()
    matched = master["amfi_code"].astype(str).isin(nav["amfi_code"].astype(str)).sum()
    page(
        "4. Data quality and validation",
        [
            _paragraph(
                f"All {matched} distinct fund-master AMFI codes are present in the raw NAV history. "
                f"The master contains {code_count} unique codes and NAV covers {nav_codes} distinct codes. "
                "AMFI scheme codes are opaque identifiers; their digits are not interpreted as decoded attributes.",
                body,
            ),
            _pdf_table(
                [
                    ["Check", "Observed result", "Handling"],
                    ["AMFI code coverage", f"{matched}/{code_count} master codes found in NAV", "Exact key coverage validated"],
                    ["Calendar-filled NAV", "Cleaned NAV includes forward-filled dates", "Returns/risk use observed market dates"],
                    ["Five-year history", "Unavailable within observed 2022-2026 span", "Five-year CAGR reported as NA"],
                    ["SIP YoY growth", "12 initial 2022 values absent", "Retain missing; do not impute"],
                    ["Category inflows", "FY25 only", "Limit category comparisons to available year"],
                    ["Holdings snapshot", "One sample snapshot; 34 schemes", "Sector HHI is snapshot-specific"],
                ],
                [1.6 * inch, 2.7 * inch, 2.5 * inch],
            ),
            Spacer(1, 0.2 * inch),
            _paragraph(
                "Forward-filled calendar NAV adds non-trading dates and can create artificial zero returns. "
                "It remains useful for continuity and dashboard display, but is excluded from the daily "
                "return-based risk calculations.",
                body,
            ),
        ],
    )

    # 6. ETL
    page(
        "5. ETL and analytical data model",
        [
            _paragraph(
                "The ETL reads raw CSV files, normalizes types and identifiers, performs data-quality checks, "
                "writes processed CSVs, and loads the project SQLite database. SQL schema and query files "
                "document the resulting relational model.",
                body,
            ),
            _pdf_table(
                [
                    ["Layer", "Responsibilities", "Key outputs"],
                    ["Raw", "Preserve received CSVs; optional live NAV utility", "data/raw/"],
                    ["Transform", "Normalize dates/codes, clean fields, validate keys", "day2_etl.py"],
                    ["Processed", "Clean analysis-ready tables", "data/processed/*.csv"],
                    ["Storage", "Relational tables and constraints", "bluestock_mf.db, schema.sql"],
                    ["Analytics", "EDA, returns/risk, investor and holdings analysis", "notebooks/, analytics/"],
                    ["Presentation", "Interactive and static dashboards/reports", "dashboard/, reports/"],
                ],
                [1.0 * inch, 3.3 * inch, 2.5 * inch],
            ),
            Spacer(1, 0.15 * inch),
            _paragraph(
                "The `amfi_code` field links scheme-level tables and `date` links dated observations. "
                "Transaction, NAV, benchmark and snapshot datasets retain their own natural grains; "
                "joining them requires explicit aggregation and date alignment.",
                body,
            ),
        ],
    )

    # 7-9 EDA
    page(
        "6. EDA: industry scale and flows",
        [
            _paragraph(
                f"The latest AUM date in the supplied fund-house snapshots is {latest_aum_date:%B %Y}; "
                f"combined reported AUM is INR {latest_aum:.2f} lakh crore. The December 2025 SIP series "
                f"ends at INR {latest_sip:,.0f} crore, its maximum observation in the supplied 48-month series.",
                body,
            ),
            _chart("05_aum_latest_ranking.png", height=2.7),
            Spacer(1, 0.12 * inch),
            _chart("06_sip_monthly_inflows.png", height=2.7),
            _paragraph("Source: supplied AUM and SIP snapshots. Amounts are reported values, not forecasts.", small),
        ],
    )
    page(
        "7. EDA: investor activity and geography",
        [
            _paragraph(
                f"The transaction sample contains {len(transactions):,} records over "
                f"{pd.to_datetime(transactions['transaction_date']).min():%b %Y} to "
                f"{pd.to_datetime(transactions['transaction_date']).max():%b %Y}. State, city-tier, "
                "age-band and transaction-type views describe these records only.",
                body,
            ),
            _chart("15_sip_amount_by_state.png", height=2.6),
            Spacer(1, 0.12 * inch),
            _chart("11_investor_age_distribution.png", height=2.6),
            _paragraph("Transaction totals are sample summaries and are not estimates of state-level market share.", small),
        ],
    )
    page(
        "8. EDA: scheme and portfolio structure",
        [
            _paragraph(
                "Scheme performance and expense comparisons show variation across the available sample. "
                "The category-flow data spans FY25 only; holdings represent one snapshot, so these figures "
                "should not be generalized to current portfolios.",
                body,
            ),
            _chart("23_performance_3yr_vs_5yr.png", height=2.65),
            Spacer(1, 0.12 * inch),
            _chart("22_sector_allocation_bar.png", height=2.65),
            _paragraph("The sector chart summarizes supplied holdings and is not an AUM-weighted industry estimate.", small),
        ],
    )

    # 10. Performance methodology
    page(
        "9. Performance methodology",
        [
            _pdf_table(
                [
                    ["Metric", "Definition / implementation"],
                    ["Daily return", "NAV(t) / NAV(t-1) - 1 on common observed market dates"],
                    ["CAGR", "(NAV_end / NAV_start)^(1/years) - 1; 1y/3y lookback, 5y unavailable"],
                    ["Sharpe", "Annualized mean daily excess return / sample SD; risk-free rate 6.5% p.a."],
                    ["Sortino", "Annualized excess return / downside deviation of negative excess-return days"],
                    ["Alpha / beta", "OLS on daily fund vs NIFTY100 returns; alpha annualized as intercept x 252"],
                    ["Max drawdown", "Minimum NAV / running peak - 1; peak/trough and recovery recorded"],
                    ["Tracking error", "Sample SD of active daily returns x sqrt(252)"],
                ],
                [1.35 * inch, 5.45 * inch],
            ),
            Spacer(1, 0.18 * inch),
            _paragraph(
                "For the return analysis, 1,150 common market dates provide 1,149 daily observations per fund. "
                "The annualization convention follows 252 trading days. Beta, alpha and tracking error remain "
                "sample- and benchmark-dependent estimates.",
                body,
            ),
        ],
    )

    # 11. Performance outputs
    leaders = performance.sort_values("score_rank").head(6)
    score_rows: list[list[object]] = [["Rank", "Scheme", "3y CAGR", "Sharpe", "Alpha (ann.)", "Score"]]
    for _, row in leaders.iterrows():
        score_rows.append(
            [
                int(row["score_rank"]),
                str(row["scheme_name"])[:42],
                _format_percent(row["cagr_3yr"]),
                f"{row['sharpe_ratio']:.2f}",
                _format_percent(row["annualized_alpha"]),
                f"{row['fund_score']:.1f}",
            ]
        )
    page(
        "10. Performance results and composite scorecard",
        [
            _paragraph(
                "The scorecard combines percentile ranks for three-year return (30%), Sharpe (25%), "
                "annualized alpha (20%), inverse expense ratio (15%), and inverse maximum drawdown (10%). "
                "It is a transparent ranking heuristic, not an expected-return model.",
                body,
            ),
            _pdf_table(score_rows, [0.42 * inch, 3.15 * inch, 0.75 * inch, 0.58 * inch, 0.93 * inch, 0.55 * inch]),
            Spacer(1, 0.2 * inch),
            _chart("benchmark_comparison_top5.png", height=3.2),
            _paragraph(
                "Top-five composite-score funds and NIFTY 50/NIFTY 100 are rebased to 100 over the latest "
                "available three-year common window. Past relative performance is not predictive.",
                small,
            ),
        ],
    )

    # 12. Risk metrics
    risk_rows: list[list[object]] = [["Fund", "Risk grade", "95% VaR (daily)", "CVaR (daily)"]]
    for _, row in var_cvar.sort_values("var_95_daily_return_pct").head(6).iterrows():
        risk_rows.append(
            [
                str(row["scheme_name"])[:44],
                str(row["risk_grade"]),
                f"{row['var_95_daily_return_pct']:.2f}%",
                f"{row['cvar_95_daily_return_pct']:.2f}%",
            ]
        )
    page(
        "11. Risk analysis: tail losses and rolling Sharpe",
        [
            _paragraph(
                f"The most negative historical 95% one-day VaR in this sample is for "
                f"{worst_var['scheme_name']}: {worst_var['var_95_daily_return_pct']:.2f}%; its "
                f"conditional tail mean is {worst_var['cvar_95_daily_return_pct']:.2f}%. "
                "Historical VaR/CVaR are distribution summaries, not loss guarantees.",
                body,
            ),
            _pdf_table(risk_rows, [3.9 * inch, 1.1 * inch, 0.9 * inch, 0.9 * inch]),
            Spacer(1, 0.12 * inch),
            _chart("rolling_sharpe_chart.png", height=2.7),
            _paragraph("Rolling Sharpe uses a 90-observation window and annualization by sqrt(252).", small),
        ],
    )

    # 13. Investor + concentration
    page(
        "12. Investor cohorts, SIP continuity and concentration",
        [
            _paragraph(
                "Investor-level analyses group the sample by first transaction year and summarize investment "
                "amounts and scheme preference. SIP continuity applies the stated eligibility of at least six "
                "SIP transactions; investors with mean inter-transaction gaps above 35 days are marked at-risk.",
                body,
            ),
            _pdf_table(
                [
                    ["Analysis", "Result", "Caveat"],
                    ["Cohorts", "First transaction year; average SIP and total invested", "Sample-period only"],
                    ["SIP continuity", "6+ SIP transactions; mean gap >35 days = at-risk", "Same-date duplicate rows can create zero-day gaps"],
                    [
                        "At-risk SIP flag",
                        f"{at_risk_sip_investors:,} of {eligible_sip_investors:,} eligible investors = {at_risk_pct:.2f}%",
                        "Rule-based flag, not churn probability",
                    ],
                    [
                        "SIP continuity rate",
                        f"{float(continuity_metrics['continuity_rate_pct']):.2f}%",
                        "Complement of the stated at-risk classification",
                    ],
                    ["Sector concentration", "HHI = sum of squared sector weights per scheme", "Single holdings snapshot; aggregation is sector-level"],
                    [
                        "Highest observed HHI",
                        f"{top_hhi['scheme_name']}: {top_hhi['hhi']:.1f}" if top_hhi is not None else "No holdings HHI available",
                        "Higher HHI means more concentrated by this measure",
                    ],
                ],
                [1.35 * inch, 2.75 * inch, 2.7 * inch],
            ),
            Spacer(1, 0.18 * inch),
            _paragraph(
                "These rules support monitoring and exploration; they do not establish investor intent or "
                "causal reasons for missed contributions.",
                body,
            ),
        ],
    )

    # 14-15 dashboard pages
    page(
        "13. Dashboard: industry and fund performance",
        [
            _paragraph(
                "The Streamlit dashboard provides four interactive sections with fund-house, category, "
                "plan and date filters where relevant. The industry page summarizes AUM, SIP inflow, "
                "folios and reported scheme counts; the performance page provides fund comparisons.",
                body,
            ),
            _chart("page_1_industry_overview.png", width=6.9, height=4.8),
            Spacer(1, 0.08 * inch),
            _chart("page_2_fund_performance.png", width=6.9, height=2.15),
        ],
    )
    page(
        "14. Dashboard: investor analytics and market trends",
        [
            _paragraph(
                "Investor analytics shows sample transaction geography, mix and demographic segments. "
                "The SIP/market page compares monthly SIP and market context and includes category flow views. "
                "Charts are limited by the coverage of their source tables.",
                body,
            ),
            _chart("page_3_investor_analytics.png", width=6.9, height=4.8),
            Spacer(1, 0.08 * inch),
            _chart("page_4_sip_and_market_trends.png", width=6.9, height=2.15),
        ],
    )

    # 16. Limits/recommendations
    page(
        "15. Limitations, recommendations and sign-off",
        [
            _pdf_table(
                [
                    ["Limitation", "Recommended next step"],
                    ["40-scheme sample; no complete market universe", "Expand scheme coverage and benchmark mapping"],
                    ["Five-year history unavailable", "Extend observed NAV history before publishing five-year CAGR"],
                    ["Category inflows FY25 only; holdings snapshot only", "Add recurring snapshots and explicit as-of dates"],
                    ["Investor sample is not population-representative", "Document sampling design and protect privacy"],
                    ["Historical metrics and arbitrary score weights", "Add sensitivity checks and suitability review before decisions"],
                    ["Dashboard is local; no PBIX/hosted URL", "Publish only after approval, access controls and data review"],
                ],
                [2.85 * inch, 3.95 * inch],
            ),
            Spacer(1, 0.2 * inch),
            _paragraph(
                "Recommendation: use the scorecard and risk outputs as screening aids, then validate holdings, "
                "benchmark fit, fees, time horizon, and investor suitability. Re-run the pipeline when source "
                "datasets are updated; keep observed and imputed NAV clearly distinguished.",
                body,
            ),
            Spacer(1, 0.15 * inch),
            _paragraph(
                "Self-review: local ETL, analytics, Streamlit dashboard, PDF report and presentation are "
                "included. Native Power BI PBIX and public hosting remain explicitly out of scope/unverified.",
                styles["center"],
            ),
        ],
    )

    output = REPORTS / "Final_Report.pdf"
    output.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(output), pagesize=letter,
        leftMargin=0.65 * inch, rightMargin=0.65 * inch,
        topMargin=0.68 * inch, bottomMargin=0.72 * inch,
        title="Bluestock Mutual Fund Analysis - Final Capstone Report",
        author="Bluestock Mutual Fund Analysis",
    )
    document.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return output


def _add_slide_title(slide: object, title: str, kicker: str = "BLUESTOCK | MUTUAL FUND ANALYSIS") -> None:
    """Add consistent title and accent line to a presentation slide."""
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(12.0), Inches(0.65))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = "Aptos Display"
    paragraph.font.size = Pt(27)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor.from_string(NAVY)
    accent = slide.shapes.add_shape(1, Inches(0.65), Inches(1.12), Inches(1.2), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor.from_string(TEAL)
    accent.line.fill.background()
    label = slide.shapes.add_textbox(Inches(10.0), Inches(0.49), Inches(2.7), Inches(0.25))
    p = label.text_frame.paragraphs[0]
    p.text = kicker
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string(BLUE)
    p.alignment = PP_ALIGN.RIGHT


def _add_bullets(slide: object, items: list[str], x: float = 0.8, y: float = 1.55, w: float = 5.4, h: float = 5.3, font_size: int = 19) -> None:
    """Place a compact bulleted text box on a slide."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor.from_string(NAVY)
        paragraph.space_after = Pt(13)


def _add_image(slide: object, path: Path, x: float, y: float, w: float, h: float) -> None:
    """Add an existing visual scaled into the requested slide box."""
    if not path.exists():
        raise FileNotFoundError(f"Required presentation image is missing: {path}")
    with PILImage.open(path) as image:
        image_width, image_height = image.size
    scale = min(w / image_width, h / image_height)
    draw_width, draw_height = image_width * scale, image_height * scale
    slide.shapes.add_picture(
        str(path),
        Inches(x + (w - draw_width) / 2),
        Inches(y + (h - draw_height) / 2),
        width=Inches(draw_width),
        height=Inches(draw_height),
    )


def _add_note(slide: object, text: str) -> None:
    """Add a small caveat or source note along the slide footer."""
    shape = slide.shapes.add_textbox(Inches(0.7), Inches(7.12), Inches(11.9), Inches(0.2))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(8)
    paragraph.font.color.rgb = RGBColor.from_string(MUTED)


def _blank_slide(presentation: Presentation, title: str) -> object:
    """Create a light-background slide with a consistent heading."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    _add_slide_title(slide, title)
    return slide


def build_presentation() -> Path:
    """Generate the requested 12-slide project presentation."""
    master = _read_data("01_fund_master.csv")
    transactions = _read_data("08_investor_transactions.csv")
    scorecard = _read_report("fund_scorecard.csv").sort_values("score_rank")
    var_cvar = _read_report("var_cvar_report.csv").sort_values("var_95_daily_return_pct")
    aum = _read_data("03_aum_by_fund_house.csv")
    sip = _read_data("04_monthly_sip_inflows.csv")
    nav = _read_data("02_nav_history.csv")
    latest_aum_date = pd.to_datetime(aum["date"]).max()
    latest_aum = aum.loc[pd.to_datetime(aum["date"]).eq(latest_aum_date), "aum_lakh_crore"].sum()
    top = scorecard.iloc[0]
    worst = var_cvar.iloc[0]

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    # 1. Title
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    title = slide.shapes.add_textbox(Inches(0.85), Inches(1.5), Inches(11.6), Inches(1.6))
    title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = title.text_frame.paragraphs[0]
    p.text = "Bluestock Mutual Fund Analysis"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string(WHITE)
    subtitle = slide.shapes.add_textbox(Inches(0.9), Inches(3.25), Inches(10.8), Inches(1.2))
    p = subtitle.text_frame.paragraphs[0]
    p.text = "Data engineering | Performance & risk | Investor analytics | Interactive dashboard"
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor.from_string("D6E4F0")
    _add_note(slide, "Final capstone presentation | Built from the supplied project datasets")

    # 2. Problem and objective
    slide = _blank_slide(presentation, "Problem and project objectives")
    _add_bullets(slide, [
        "Bring fragmented industry, scheme, NAV, benchmark and transaction data into one analysis workflow.",
        "Compare returns and risk consistently across a 40-scheme sample.",
        "Explore investor transaction patterns and SIP continuity without claiming population representativeness.",
        "Deliver reproducible local analytics and an interactive four-page Streamlit dashboard.",
    ], w=11.5, font_size=20)

    # 3. Data sources
    slide = _blank_slide(presentation, "Data sources and coverage")
    _add_bullets(slide, [
        "10 supplied CSV datasets: fund master, NAV, AUM, SIP, category flows, folios, performance, transactions, holdings and benchmarks.",
        f"NAV: {nav['amfi_code'].nunique()} schemes across {pd.to_datetime(nav['date']).nunique():,} observed dates.",
        f"Investor transaction sample: {len(transactions):,} rows, January 2024 to May 2025.",
        f"Latest fund-house AUM snapshot: {latest_aum_date:%b %Y}; INR {latest_aum:.2f} lakh crore.",
        "Additional MFAPI live-NAV utility is opt-in; analytics pipeline runs offline by default.",
    ], w=11.6, font_size=18)
    _add_note(slide, "Coverage windows differ by dataset; category inflows are FY25-only and holdings are a single snapshot.")

    # 4. Architecture
    slide = _blank_slide(presentation, "ETL and analytics architecture")
    stages = [
        ("RAW DATA", "10 source CSVs"),
        ("TRANSFORM", "Validate + normalize"),
        ("MODEL", "Processed CSV + SQLite"),
        ("ANALYZE", "EDA + performance + risk"),
        ("DELIVER", "Streamlit + PDF + PPTX"),
    ]
    for index, (label, caption) in enumerate(stages):
        x = 0.55 + index * 2.57
        shape = slide.shapes.add_shape(1, Inches(x), Inches(2.4), Inches(2.18), Inches(1.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(PALE)
        shape.line.color.rgb = RGBColor.from_string(BLUE)
        text = slide.shapes.add_textbox(Inches(x + 0.12), Inches(2.68), Inches(1.95), Inches(1.0))
        frame = text.text_frame
        frame.paragraphs[0].text = label
        frame.paragraphs[0].font.bold = True
        frame.paragraphs[0].font.size = Pt(15)
        frame.paragraphs[0].font.color.rgb = RGBColor.from_string(BLUE)
        p = frame.add_paragraph()
        p.text = caption
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor.from_string(NAVY)
        if index < len(stages) - 1:
            arrow = slide.shapes.add_textbox(Inches(x + 2.18), Inches(2.82), Inches(0.38), Inches(0.4))
            arrow.text_frame.paragraphs[0].text = ">"
            arrow.text_frame.paragraphs[0].font.size = Pt(21)
            arrow.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(TEAL)
    _add_bullets(slide, ["Keys: amfi_code for scheme-level joins; date for aligned snapshots/market series."], y=4.7, w=11.3, h=1.5, font_size=17)

    # 5-6. EDA highlights
    slide = _blank_slide(presentation, "EDA highlight 1 | Industry and flows")
    _add_image(slide, CHARTS / "05_aum_latest_ranking.png", 0.6, 1.5, 5.95, 4.75)
    _add_image(slide, CHARTS / "06_sip_monthly_inflows.png", 6.8, 1.5, 5.95, 4.75)
    _add_note(slide, f"December 2025 SIP inflow: INR {float(sip.iloc[-1]['sip_inflow_crore']):,.0f} crore; supplied AUM snapshots through {latest_aum_date:%b %Y}.")

    slide = _blank_slide(presentation, "EDA highlight 2 | Investors and schemes")
    _add_image(slide, CHARTS / "15_sip_amount_by_state.png", 0.6, 1.5, 5.95, 4.75)
    _add_image(slide, CHARTS / "23_performance_3yr_vs_5yr.png", 6.8, 1.5, 5.95, 4.75)
    _add_note(slide, "Investor summaries describe the supplied sample. Five-year returns are not available for the complete 2022-2026 NAV window.")

    # 7. Performance
    slide = _blank_slide(presentation, "Performance | Composite scorecard")
    _add_bullets(slide, [
        f"Top composite rank: {top['scheme_name']}",
        f"Score {top['fund_score']:.2f}/100 | three-year CAGR {_format_percent(top['cagr_3yr'])} | Sharpe {top['sharpe_ratio']:.2f}",
        "Weights: 30% three-year return, 25% Sharpe, 20% annualized alpha, 15% inverse expense, 10% inverse drawdown.",
        "Five-year CAGR is NA: observed daily NAV begins in January 2022.",
    ], x=0.8, w=6.0, font_size=17)
    _add_image(slide, REPORTS / "benchmark_comparison_top5.png", 7.0, 1.65, 5.6, 4.55)
    _add_note(slide, "Score is a screening heuristic using percentile ranks; not an investment recommendation.")

    # 8. Risk
    slide = _blank_slide(presentation, "Performance | Risk and benchmark context")
    _add_bullets(slide, [
        f"Most negative historical 95% daily VaR: {worst['scheme_name']} ({worst['var_95_daily_return_pct']:.2f}%).",
        f"Tail conditional mean (CVaR): {worst['cvar_95_daily_return_pct']:.2f}%.",
        "Daily metrics use 1,149 observed market-date returns per scheme; annualized with 252 trading days.",
        "Risk-free proxy is 6.5% p.a.; alpha/beta use OLS versus NIFTY100 daily returns.",
    ], w=6.0, font_size=16)
    _add_image(slide, REPORTS / "rolling_sharpe_chart.png", 6.95, 1.65, 5.6, 4.65)
    _add_note(slide, "Historical tail estimates describe the sample and do not bound future losses.")

    # 9-10. Dashboard screenshots
    slide = _blank_slide(presentation, "Dashboard | Industry overview and performance")
    _add_image(slide, PAGES / "page_1_industry_overview.png", 0.55, 1.45, 6.1, 5.25)
    _add_image(slide, PAGES / "page_2_fund_performance.png", 6.8, 1.45, 6.0, 5.25)
    _add_note(slide, "Interactive local Streamlit dashboard; screenshots show static page exports.")

    slide = _blank_slide(presentation, "Dashboard | Investor and SIP trends")
    _add_image(slide, PAGES / "page_3_investor_analytics.png", 0.55, 1.45, 6.1, 5.25)
    _add_image(slide, PAGES / "page_4_sip_and_market_trends.png", 6.8, 1.45, 6.0, 5.25)
    _add_note(slide, "Page outputs are limited to source dataset coverage and are not a published dashboard.")

    # 11. Findings and recommendations
    slide = _blank_slide(presentation, "Key findings and recommendations")
    _add_bullets(slide, [
        "Industry AUM totals INR {:.2f} lakh crore in the latest supplied snapshot; monthly SIP series peaks at INR {:,.0f} crore.".format(latest_aum, float(sip["sip_inflow_crore"].max())),
        "Observed market-date NAV avoids artificial zero returns from forward-filled weekends and holidays.",
        "At-risk SIP flag is a rule-based signal: mean gap >35 days for investors with 6+ SIP rows.",
        "Use scorecard and risk metrics for screening; verify suitability, holdings, fees and benchmark fit.",
        "Extend category flow, holdings and NAV history before stronger claims or long-horizon comparisons.",
    ], w=11.5, font_size=17)

    # 12. Thank you
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    text = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.8), Inches(1.1))
    p = text.text_frame.paragraphs[0]
    p.text = "Thank you"
    p.font.size = Pt(42)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string(WHITE)
    details = slide.shapes.add_textbox(Inches(0.85), Inches(3.35), Inches(11.6), Inches(1.3))
    p = details.text_frame.paragraphs[0]
    p.text = "Questions | Local Streamlit dashboard | Reproducible analytics"
    p.font.size = Pt(21)
    p.font.color.rgb = RGBColor.from_string("D6E4F0")
    _add_note(slide, "Limitations: sample coverage, historical metrics, no hosted URL and no validated native PBIX.")

    output = REPORTS / "Bluestock_MF_Presentation.pptx"
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    if len(presentation.slides) != 12:
        raise RuntimeError(f"Presentation should have 12 slides; found {len(presentation.slides)}")
    return output


def build() -> tuple[Path, Path]:
    """Create both final deliverables and return their paths."""
    return build_pdf(), build_presentation()


if __name__ == "__main__":
    for deliverable in build():
        print(f"Created {deliverable.relative_to(ROOT)}")
